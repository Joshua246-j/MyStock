import pymupdf as fitz
import pytest
from docx import Document

from app.documents.loader import DocumentExtractionError
from app.documents.pdf import PDFLoader
from app.documents.text import HTMLLoader, TextLoader
from app.documents.word import WordLoader


def test_txt_loader_requires_utf8_and_preserves_text(tmp_path):
    path = tmp_path / "research.txt"
    path.write_text("Revenue increased 12 percent.", encoding="utf-8")
    result = TextLoader().extract_text(str(path), path.name, 9)
    assert result[0]["text"] == "Revenue increased 12 percent."
    assert result[0]["metadata"]["document_id"] == 9
    path.write_bytes(b"\xff\xfe")
    with pytest.raises(DocumentExtractionError, match="UTF-8"):
        TextLoader().extract_text(str(path), path.name, 9)


def test_html_loader_extracts_title_and_discards_script_content(tmp_path):
    path = tmp_path / "report.html"
    path.write_text("<html><head><title>Results</title><script>ignore me</script></head><body><h1>Revenue</h1><p>Revenue was 12.</p></body></html>", encoding="utf-8")
    result = HTMLLoader().extract_text(str(path), path.name, 10)
    assert result[0]["metadata"]["title"] == "Results"
    assert "Revenue was 12." in result[0]["text"]
    assert "ignore me" not in result[0]["text"]


def test_scanned_pdf_without_tesseract_reports_ocr_requirement(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    document = fitz.open()
    page = document.new_page()
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 80, 80), 0)
    pixmap.clear_with(255)
    page.insert_image(fitz.Rect(20, 20, 100, 100), stream=pixmap.tobytes("png"))
    document.save(path)
    document.close()
    monkeypatch.setattr("app.documents.pdf.shutil.which", lambda _name: None)
    with pytest.raises(DocumentExtractionError, match="appears to be scanned.*Tesseract"):
        PDFLoader().extract_text(str(path), path.name, 11)


def test_pdf_loader_preserves_page_numbers(tmp_path):
    path = tmp_path / "filing.pdf"
    document = fitz.open()
    document.new_page().insert_text((72, 72), "Quarterly revenue was 125 million dollars.")
    document.new_page().insert_text((72, 72), "Operating costs increased modestly.")
    document.save(path)
    document.close()
    result = PDFLoader().extract_text(str(path), path.name, 12)
    assert [item["metadata"]["page"] for item in result] == [1, 2]


def test_word_loader_preserves_heading_sections_and_tables(tmp_path):
    path = tmp_path / "filing.docx"
    document = Document()
    document.add_heading("Risk Factors", level=1)
    document.add_paragraph("Supply disruption could affect revenue.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Fiscal year"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "2026"
    table.cell(1, 1).text = "$735 million"
    document.save(path)

    result = WordLoader().extract_text(str(path), path.name, 13)
    assert result[0]["metadata"]["section"] == "Risk Factors"
    assert result[1]["metadata"]["section"] == "Risk Factors"
    assert "735 million" in result[2]["text"]
    assert result[2]["metadata"]["document_type"] == "docx"
